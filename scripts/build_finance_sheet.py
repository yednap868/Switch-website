#!/usr/bin/env python3
"""Build a monthly finance workbook for Switch: Earnings + Expenses + auto Summary."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MONTH = "July 2026"
OUT = "/Users/alt/Downloads/Switch_Finance_July_2026.xlsx"

# --- palette ---
DARK = "1A1A1A"      # header bg
WHITE = "FFFFFF"
GREEN = "1B7F4B"     # earnings
RED = "B23A3A"       # expenses
LIGHT = "F2F2F2"     # zebra / label bg
ACCENT = "EAF3EE"

thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row, ncols, bg=DARK):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(bold=True, color=WHITE, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

def make_ledger(ws, title, color, categories):
    ws.sheet_view.showGridLines = False
    # Title
    ws.merge_cells("A1:E1")
    t = ws["A1"]
    t.value = f"Switch — {title} · {MONTH}"
    t.font = Font(bold=True, size=15, color=WHITE)
    t.fill = PatternFill("solid", fgColor=color)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    headers = ["Date", "Party / Description", "Category", "Notes", "Amount (₹)"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=2, column=i, value=h)
    style_header(ws, 2, 5)

    # 40 blank data rows with borders + zebra
    first, last = 3, 42
    for r in range(first, last + 1):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(row=r, column=1).number_format = "dd-mmm-yyyy"
        ws.cell(row=r, column=5).number_format = '#,##0'
    # category dropdown
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(categories), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C{first}:C{last}")

    # total row
    tr = last + 1
    ws.cell(row=tr, column=4, value="TOTAL").font = Font(bold=True, size=12)
    ws.cell(row=tr, column=4).alignment = Alignment(horizontal="right")
    tot = ws.cell(row=tr, column=5, value=f"=SUM(E{first}:E{last})")
    tot.font = Font(bold=True, size=12, color=WHITE)
    tot.fill = PatternFill("solid", fgColor=color)
    tot.number_format = '#,##0'
    tot.border = border

    widths = [14, 40, 22, 30, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return f"'{ws.title}'!E{tr}"  # ref to total cell

wb = openpyxl.Workbook()

# --- Earnings ---
ws_e = wb.active
ws_e.title = "Earnings"
earn_cats = ["Client Staffing Fee", "Placement Commission", "Subscription / Retainer",
             "Onboarding Fee", "Other Income"]
earn_total_ref = make_ledger(ws_e, "Earnings", GREEN, earn_cats)

# --- Expenses ---
ws_x = wb.create_sheet("Expenses")
exp_cats = ["Worker Payouts", "Salaries & Stipends", "Office / Rent", "Marketing",
            "Software & Tools", "Travel", "Legal & Compliance", "Misc"]
exp_total_ref = make_ledger(ws_x, "Expenses", RED, exp_cats)

# --- Summary (first in tab order) ---
ws_s = wb.create_sheet("Summary", 0)
ws_s.sheet_view.showGridLines = False
ws_s.merge_cells("A1:C1")
s = ws_s["A1"]
s.value = f"Switch — Finance Summary · {MONTH}"
s.font = Font(bold=True, size=16, color=WHITE)
s.fill = PatternFill("solid", fgColor=DARK)
s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws_s.row_dimensions[1].height = 34

rows = [
    ("Total Earnings", f"={earn_total_ref}", GREEN),
    ("Total Expenses", f"={exp_total_ref}", RED),
    ("Net Profit / (Loss)", f"={earn_total_ref}-{exp_total_ref}", DARK),
]
r = 3
for label, formula, color in rows:
    lc = ws_s.cell(row=r, column=1, value=label)
    lc.font = Font(bold=True, size=12)
    lc.fill = PatternFill("solid", fgColor=ACCENT)
    lc.alignment = Alignment(vertical="center", indent=1)
    lc.border = border
    ws_s.cell(row=r, column=2).border = border  # spacer col styling
    vc = ws_s.cell(row=r, column=3, value=formula)
    vc.font = Font(bold=True, size=12, color=WHITE)
    vc.fill = PatternFill("solid", fgColor=color)
    vc.number_format = '₹ #,##0'
    vc.alignment = Alignment(horizontal="right", indent=1)
    vc.border = border
    ws_s.row_dimensions[r].height = 26
    r += 1

# margin % row
mc = ws_s.cell(row=r, column=1, value="Profit Margin")
mc.font = Font(bold=True, size=12)
mc.fill = PatternFill("solid", fgColor=ACCENT)
mc.alignment = Alignment(vertical="center", indent=1)
mc.border = border
ws_s.cell(row=r, column=2).border = border
pm = ws_s.cell(row=r, column=3, value=f"=IF({earn_total_ref}=0,0,({earn_total_ref}-{exp_total_ref})/{earn_total_ref})")
pm.number_format = '0.0%'
pm.font = Font(bold=True, size=12)
pm.alignment = Alignment(horizontal="right", indent=1)
pm.border = border

ws_s.column_dimensions["A"].width = 26
ws_s.column_dimensions["B"].width = 4
ws_s.column_dimensions["C"].width = 20

note = ws_s.cell(row=r + 2, column=1,
    value="Fill in the Earnings and Expenses tabs — totals here update automatically.")
note.font = Font(italic=True, size=10, color="666666")
ws_s.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=3)

wb.save(OUT)
print("Saved:", OUT)
print("Tabs:", wb.sheetnames)
