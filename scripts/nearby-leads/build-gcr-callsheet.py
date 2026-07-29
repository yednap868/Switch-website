#!/usr/bin/env python3
"""Golf Course Road corridor calling workbook — one tab per sector.

Sectors (in the user's order): 56, 55, 53, 52A, 42, 27, 28, 26A.
Source: gcr-south.csv + gcr-north.csv (Google Places crawls of the corridor).
Each business is bucketed into a sector by the "Sector NN" in its Google address.

Only businesses WITH a phone number are kept, deduped by number across the whole
workbook. A "Number Type" column flags Mobile (likely owner/manager) vs Landline
(reception); mobiles are sorted to the top so callers reach decision-makers first.

Tabs: Performance (live dashboard) · one per sector · How to use.
Opens in Excel; imports into Google Sheets with dropdowns + formulas intact.

    python3 build-gcr-callsheet.py
"""
import csv, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = "switch-golf-course-road-callsheet.xlsx"
SRC_FILES = ["gcr-south.csv", "gcr-north.csv"]

# Sector -> regex that matches its label in a Google formatted address.
# Order here = tab order = bucket priority (first match wins).
SECTORS = [
    ("Sector 56",  r"sector[\s\-]*56\b"),
    ("Sector 55",  r"sector[\s\-]*55\b"),
    ("Sector 53",  r"sector[\s\-]*53\b"),
    ("Sector 52A", r"sector[\s\-]*52[\s\-]*a\b"),
    ("Sector 42",  r"sector[\s\-]*42\b"),
    ("Sector 27",  r"sector[\s\-]*27\b"),
    ("Sector 28",  r"sector[\s\-]*28\b"),
    ("Sector 26A", r"sector[\s\-]*26[\s\-]*a\b"),
]

# ── palette ─────────────────────────────────────────────────────────────────
INDIGO="4F46E5"; INDIGO_L="EEF0FB"; DARK="111827"; GREY="6B7280"
GREEN="C6EFCE"; GREEN_TX="1E7E34"; RED="FFC7CE"; AMBER="FFF3CD"
WHITE="FFFFFF"; ZEBRA="F7F7FB"; BLUE_TX="1D4ED8"
thin = Side(style="thin", color="E2E2EC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
def fill(c): return PatternFill("solid", fgColor=c)

# ── roles Switch can staff, by category/name ────────────────────────────────
def roles_for(cat, name):
    t = ("%s %s" % (cat, name)).lower()
    def has(*kws): return any(k in t for k in kws)
    if has("warehouse","logistic","cargo","courier","freight","storage","godown",
           "packers","movers","transport","distribution","supply","3pl","fulfil","shipping"):
        return "Loader, Packer, Picker, Warehouse Helper, Delivery"
    if has("restaurant","cafe","coffee","food","kitchen","bakery","fast_food","meal","dhaba","sweet","bar","pizza"):
        return "Kitchen Helper, Cook, Dishwasher, Waiter"
    if has("hotel","lodging","guest_house","hostel","motel","resort","banquet"):
        return "Housekeeping, Room Boy, Kitchen Helper, Front Desk"
    if has("hospital","clinic","medical","dental","dentist","health","pharmacy","doctor","diagnostic","physio"):
        return "Ward Boy, Housekeeping, Attendant, Security"
    if has("mall","supermarket","store","retail","clothing","grocery","general_store","shop","market","boutique"):
        return "Store Helper, Housekeeping, Security, Sales Staff"
    if has("salon","beauty","spa","hairdresser","massage","parlour","nail","wellness"):
        return "Beautician, Helper, Receptionist, Cleaner"
    if has("gym","fitness","sports","yoga"):
        return "Trainer, Housekeeping, Front Desk"
    if has("fuel","petrol","cng","charging"):
        return "Pump Attendant, Cleaner, Security"
    if has("manufacturer","factory","industrial","industries","engineering"):
        return "Machine Operator, Helper, Loader, Packer"
    if has("car_repair","garage","motor","automobile","service_station"):
        return "Mechanic Helper, Cleaner, Washer"
    if has("school","educational","institute","college","coaching","university","preschool","daycare"):
        return "Peon, Housekeeping, Security, Attendant"
    if has("bank","office","corporate","consultant","association","finance","insurance","real_estate","government","coworking"):
        return "Office Boy, Housekeeping, Security, Data Entry"
    return "Helper, Housekeeping, Security, Office Boy"

# ── phone classification ────────────────────────────────────────────────────
def classify_phone(raw):
    d = re.sub(r"\D", "", raw or "")
    if d.startswith("91") and len(d) == 12:
        d = d[2:]
    if d.startswith("0") and len(d) == 11:
        d = d[1:]
    if d.startswith(("1800","1860")):
        return d, "Landline (helpline)"
    if len(d) == 10 and d[0] in "6789":
        return d, "Mobile (owner/mgr)"
    return d, "Landline (reception)"

def bucket(address):
    a = (address or "").lower()
    for name, pat in SECTORS:
        if re.search(pat, a):
            return name
    return None

def load():
    by_sector = {name: [] for name, _ in SECTORS}
    seen = set()
    for path in SRC_FILES:
        try:
            with open(path, newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    phone = (row.get("phone") or "").strip()
                    if not phone:
                        continue
                    digits, ptype = classify_phone(phone)
                    if len(digits) < 7 or digits in seen:
                        continue
                    sector = bucket(row.get("address"))
                    if not sector:
                        continue
                    seen.add(digits)
                    by_sector[sector].append({
                        "name": (row.get("name") or "").strip(),
                        "category": (row.get("primary_type") or "").strip() or "business",
                        "phone": phone,
                        "ptype": ptype,
                        "rating": (row.get("rating") or "").strip(),
                        "address": (row.get("address") or "").strip(),
                        "map": (row.get("maps_url") or "").strip(),
                    })
        except FileNotFoundError:
            print("!! missing %s" % path)
    # Mobiles first (owner/manager), then highest rating, then name.
    for name in by_sector:
        by_sector[name].sort(key=lambda r: (
            0 if r["ptype"].startswith("Mobile") else 1,
            -(float(r["rating"]) if r["rating"] else 0),
            r["name"].lower(),
        ))
    return by_sector

# ── columns ─────────────────────────────────────────────────────────────────
COLS = [
    ("S.No", 6), ("Business Name", 30), ("Category", 17),
    ("Suggested Roles (per shift)", 32), ("Phone", 15),
    ("Number Type", 18), ("Rating", 7), ("Address", 36),
    ("Interested?", 12), ("Hours (1-12)", 11), ("Order Confirmed?", 15),
    ("Remarks", 30), ("Map", 6),
]
HDR = 4
C_INT, C_HRS, C_ORD, C_MAP = "I", "J", "K", "M"

def build_sheet(wb, first, tab, rows):
    ws = wb.active if first else wb.create_sheet()
    ws.title = tab
    ws.sheet_view.showGridLines = False
    last = HDR + max(len(rows), 1)
    mobiles = sum(1 for r in rows if r["ptype"].startswith("Mobile"))

    for c,(t,w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.merge_cells("A1:M1")
    b = ws["A1"]; b.value = "SWITCH  ·  %s  ·  CALL SHEET" % tab.upper()
    b.fill = fill(INDIGO); b.font = Font(bold=True, color=WHITE, size=16)
    b.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    ws.merge_cells("A2:M2")
    s = ws["A2"]
    s.value = ("Shifts 1–12 hrs, billed per hour · pay after worker reports · instant "
               "replacement   |   %d leads · %d mobiles (owner/mgr) at top" % (len(rows), mobiles))
    s.font = Font(color=GREY, size=10, italic=True); s.alignment = Alignment(horizontal="center")

    # inline counters row 3
    def rng(col): return "%s%d:%s%d" % (col, HDR+1, col, last)
    for cell, label, formula, color in [
        ("A3","LEADS","=COUNTA(%s)"%rng("B"),DARK),
        ("C3","CONTACTED","=COUNTA(%s)"%rng(C_INT),INDIGO),
        ("E3","INTERESTED",'=COUNTIF(%s,"Yes")'%rng(C_INT),GREEN_TX),
        ("G3","ORDERS",'=COUNTIF(%s,"Yes")'%rng(C_ORD),GREEN_TX),
        ("I3","HOURS BOOKED","=SUM(%s)"%rng(C_HRS),"B7791F"),
    ]:
        col = cell[0]
        lc = ws[cell]; lc.value = label
        lc.font = Font(bold=True, color=GREY, size=8); lc.alignment = Alignment(horizontal="center")
        vc = ws["%s3" % chr(ord(col)+1)]; vc.value = formula
        vc.font = Font(bold=True, color=color, size=13); vc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 20

    for c,(title,_) in enumerate(COLS, start=1):
        cell = ws.cell(row=HDR, column=c, value=title)
        cell.fill = fill(INDIGO); cell.font = Font(bold=True, color=WHITE, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[HDR].height = 30

    WRAP = Alignment(vertical="top", wrap_text=True)
    CTR  = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for i, row in enumerate(rows):
        r = HDR + 1 + i
        vals = [i+1, row["name"], row["category"], roles_for(row["category"], row["name"]),
                row["phone"], row["ptype"], row["rating"], row["address"], "", "", "", "", ""]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = CTR if c in (1,5,6,7,9,10,11,13) else WRAP
            cell.border = BORDER
            if r % 2 == 0 and c not in (9,10,11):
                cell.fill = fill(ZEBRA)
        # tint the Number Type cell green when it's a mobile (owner/manager)
        tc = ws.cell(row=r, column=6)
        if row["ptype"].startswith("Mobile"):
            tc.font = Font(bold=True, color=GREEN_TX)
        else:
            tc.font = Font(color=GREY)
        if row["map"]:
            m = ws.cell(row=r, column=13, value="Map")
            m.hyperlink = row["map"]; m.font = Font(color="0563C1", underline="single")

    def dropdown(col, options):
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(options),
                            allow_blank=True, showDropDown=False)
        dv.prompt = "Click the arrow and choose"; dv.promptTitle = "Pick one"
        ws.add_data_validation(dv)
        dv.add("%s%d:%s%d" % (col, HDR+1, col, last))
    dropdown(C_INT, ["Yes","No","Maybe","Call back","No answer","Wrong number"])
    dropdown(C_HRS, [str(i) for i in range(1,13)])
    dropdown(C_ORD, ["Yes","No","Pending"])

    ws.conditional_formatting.add("%s%d:%s%d"%(C_INT,HDR+1,C_INT,last),
        FormulaRule(formula=['%s%d="Yes"'%(C_INT,HDR+1)], fill=fill(GREEN), font=Font(color=GREEN_TX, bold=True)))
    ws.conditional_formatting.add("%s%d:%s%d"%(C_INT,HDR+1,C_INT,last),
        FormulaRule(formula=['%s%d="Call back"'%(C_INT,HDR+1)], fill=fill(AMBER)))
    ws.conditional_formatting.add("%s%d:%s%d"%(C_ORD,HDR+1,C_ORD,last),
        FormulaRule(formula=['%s%d="Yes"'%(C_ORD,HDR+1)], fill=fill(GREEN), font=Font(color=GREEN_TX, bold=True)))
    ws.conditional_formatting.add("A%d:M%d"%(HDR+1,last),
        FormulaRule(formula=['$%s%d="Yes"'%(C_ORD,HDR+1)], fill=fill("EAF8EE")))

    ws.freeze_panes = "A%d" % (HDR+1)
    ws.auto_filter.ref = "A%d:M%d" % (HDR, last)
    return len(rows), mobiles

# ── performance tab ─────────────────────────────────────────────────────────
def build_performance(wb, areas):
    ws = wb.create_sheet("Performance")
    ws.sheet_view.showGridLines = False
    for col, w in (("A",3),("B",22),("C",12),("D",12),("E",12),("F",12),("G",12),("H",12),("I",3)):
        ws.column_dimensions[col].width = w

    def ref(tab, col, count): return "'%s'!%s%d:%s%d" % (tab, col, HDR+1, col, HDR+max(count,1))
    def leads(t,n):     return "COUNTA(%s)"        % ref(t,"B",n)
    def contacted(t,n): return "COUNTA(%s)"        % ref(t,C_INT,n)
    def interested(t,n):return 'COUNTIF(%s,"Yes")' % ref(t,C_INT,n)
    def orders(t,n):    return 'COUNTIF(%s,"Yes")' % ref(t,C_ORD,n)
    def hours(t,n):     return "SUM(%s)"           % ref(t,C_HRS,n)
    def mobiles(t,n):   return 'COUNTIF(%s,"Mobile*")' % ref(t,"F",n)
    def joinsum(fn):    return "+".join(fn(t,n) for t,n in areas)

    ws.merge_cells("B2:H3")
    b = ws["B2"]; b.value = "SWITCH  ·  GOLF COURSE ROAD  ·  CALLING PERFORMANCE"
    b.fill = fill(INDIGO); b.font = Font(bold=True, color=WHITE, size=16)
    b.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("B4:H4")
    s = ws["B4"]; s.value = "Live totals across all sector tabs — updates as you fill the call sheets"
    s.font = Font(color=GREY, size=10, italic=True); s.alignment = Alignment(horizontal="center")

    def card(cell, label, formula, color=INDIGO, fmt="0", big=True):
        col = cell[0]; row = int(cell[1:])
        lc = ws["%s%d"%(col,row)]; lc.value = label
        lc.font = Font(bold=True, color=GREY, size=9); lc.alignment = Alignment(horizontal="center")
        vc = ws["%s%d"%(col,row+1)]; vc.value = formula; vc.number_format = fmt
        vc.font = Font(bold=True, color=color, size=22 if big else 13)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row+1].height = 30
        for rr in (row,row+1): ws["%s%d"%(col,rr)].fill = fill(INDIGO_L if big else WHITE)

    card("B6","TOTAL LEADS","=%s"%joinsum(leads), color=DARK)
    card("D6","MOBILES (OWNER/MGR)","=%s"%joinsum(mobiles), color=BLUE_TX)
    card("F6","ORDERS","=%s"%joinsum(orders), color=GREEN_TX)
    card("B9","CONTACTED","=%s"%joinsum(contacted), color=INDIGO, big=False)
    card("D9","INTERESTED","=%s"%joinsum(interested), color=GREEN_TX, big=False)
    card("F9","CONVERSION","=IFERROR((%s)/(%s),0)"%(joinsum(orders),joinsum(contacted)),
         color=INDIGO, fmt="0%", big=False)
    card("H9","HOURS","=%s"%joinsum(hours), color="B7791F", big=False)

    hrow = 12
    heads = ["Sector","Leads","Mobiles","Contacted","Interested","Orders","Hours"]
    for c,h in enumerate(heads, start=2):
        cell = ws.cell(row=hrow, column=c, value=h)
        cell.fill = fill(INDIGO); cell.font = Font(bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center"); cell.border = BORDER
    for i,(t,n) in enumerate(areas):
        r = hrow+1+i
        cells = [t, "=%s"%leads(t,n), "=%s"%mobiles(t,n), "=%s"%contacted(t,n),
                 "=%s"%interested(t,n), "=%s"%orders(t,n), "=%s"%hours(t,n)]
        for c,v in enumerate(cells, start=2):
            cell = ws.cell(row=r, column=c, value=v); cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if c==2 else "center")
            if r % 2: cell.fill = fill(ZEBRA)
    tr = hrow+1+len(areas)
    totals = ["TOTAL","=%s"%joinsum(leads),"=%s"%joinsum(mobiles),"=%s"%joinsum(contacted),
              "=%s"%joinsum(interested),"=%s"%joinsum(orders),"=%s"%joinsum(hours)]
    for c,v in enumerate(totals, start=2):
        cell = ws.cell(row=tr, column=c, value=v); cell.border = BORDER; cell.fill = fill(INDIGO_L)
        cell.font = Font(bold=True, color=DARK)
        cell.alignment = Alignment(horizontal="left" if c==2 else "center")
    note = ws.cell(row=tr+2, column=2,
        value="↳ Fill Interested? / Hours / Order Confirmed? on each sector tab — these numbers move automatically.")
    note.font = Font(color=GREY, size=10, italic=True)

# ── how-to tab ──────────────────────────────────────────────────────────────
def build_howto(wb):
    how = wb.create_sheet("How to use")
    how.sheet_view.showGridLines = False
    how.column_dimensions["A"].width = 3
    how.column_dimensions["B"].width = 98
    guide = [
        ("SWITCH — GOLF COURSE ROAD CALLING SHEET  ·  HOW TO USE", "title"), ("",""),
        ("THE OFFER (say this, every time)", "head"),
        ("• Shifts of 1–12 hours, billed by the hour. NO monthly, NO weekly.", "body"),
        ("• Pay only AFTER the worker reports — no advance.", "body"),
        ("• Instant replacement if the worker doesn't show or isn't a fit.", "body"),
        ("• Verified, local staff delivered in 24 hours.", "body"), ("",""),
        ("ABOUT THE NUMBERS", "head"),
        ("• 'Number Type' = Mobile (owner/mgr)  → a personal mobile; usually the owner or manager. Call these FIRST.", "body"),
        ("• 'Number Type' = Landline (reception) → front desk; ask for the owner/manager by name.", "body"),
        ("• Mobiles are already sorted to the top of every sector tab.", "body"), ("",""),
        ("HOW TO WORK THE SHEET", "head"),
        ("1. One tab per sector: 56, 55, 53, 52A, 42, 27, 28, 26A. Work each top-down.", "body"),
        ("2. WhatsApp first, then call. After each contact, click the dropdowns:", "body"),
        ("     · Interested? → Yes / No / Maybe / Call back / No answer / Wrong number", "body"),
        ("     · Hours → 1 to 12 (what they want per shift)", "body"),
        ("     · Order Confirmed? → Yes once they agree to take staff", "body"),
        ("3. 'Suggested Roles' is a starting pitch — confirm what they actually need.", "body"),
        ("4. The Performance tab totals every sector automatically as you fill them in.", "body"), ("",""),
        ("Only businesses with a phone number are included. Deduped by number across all sectors.", "muted"),
    ]
    r = 1
    for text, kind in guide:
        cell = how.cell(row=r, column=2, value=text)
        if kind == "title":
            cell.font = Font(bold=True, color=WHITE, size=15); cell.fill = fill(INDIGO)
            cell.alignment = Alignment(horizontal="center", vertical="center"); how.row_dimensions[r].height = 30
        elif kind == "head":
            cell.font = Font(bold=True, color=INDIGO, size=12)
        elif kind == "muted":
            cell.font = Font(italic=True, color=GREY, size=10)
        else:
            cell.font = Font(color=DARK, size=11); cell.alignment = Alignment(wrap_text=True)
        r += 1

# ── main ────────────────────────────────────────────────────────────────────
wb = Workbook()
data = load()
areas = []
first = True
for name, _ in SECTORS:
    rows = data[name]
    n, mob = build_sheet(wb, first, name, rows)
    areas.append((name, n))
    print("  %-11s %4d leads  (%d mobile / owner-mgr)" % (name, n, mob))
    first = False
build_performance(wb, areas)
build_howto(wb)
perf = wb["Performance"]; wb._sheets.remove(perf); wb._sheets.insert(0, perf)
wb.active = 0
wb.save(OUT)
print("Wrote %s" % OUT)
