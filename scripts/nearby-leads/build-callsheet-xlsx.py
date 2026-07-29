#!/usr/bin/env python3
"""Build the polished, clickable "2-Day Order Blitz" workbook from the CSV.

Three tabs:
  • Dashboard   — live counters (orders vs goal, contacted, interested, hours)
  • Call Sheet  — leads with click-to-select dropdowns + colour-coded status
  • How to use  — the offer rules and a 30-second guide

Opens in Excel; imports into Google Sheets with dropdowns + formulas intact.

    python3 build-callsheet-xlsx.py        # reads 2-day-blitz-callsheet.csv
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

SRC = "2-day-blitz-callsheet.csv"
OUT = "2-day-blitz-callsheet.xlsx"
GOAL = 10

# ── palette (matches the site) ─────────────────────────────────────────────
INDIGO   = "4F46E5"
INDIGO_L = "EEF0FB"
DARK     = "111827"
GREY     = "6B7280"
GREEN    = "C6EFCE"
GREEN_TX = "1E7E34"
RED      = "FFC7CE"
AMBER    = "FFF3CD"
WHITE    = "FFFFFF"
ZEBRA    = "F7F7FB"

thin = Side(style="thin", color="E2E2EC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def fill(c): return PatternFill("solid", fgColor=c)

# ── columns ─────────────────────────────────────────────────────────────────
COLS = [
    ("Priority", 9),
    ("Business Name", 32),
    ("Category", 19),
    ("Suggested Roles (per shift)", 36),
    ("Phone", 16),
    ("Address", 40),
    ("Interested?", 15),
    ("Hours Needed (1-12)", 16),
    ("Order Confirmed?", 16),
    ("Remarks", 38),
    ("Map", 7),
]

with open(SRC, newline="", encoding="utf-8") as f:
    rows = list(csv.DictReader(f))
LAST = len(rows) + 1            # last data row on Call Sheet
SHEET = "Call Sheet"
def cs(rng): return "'%s'!%s" % (SHEET, rng)

wb = Workbook()

# ════════════════════════════ 1. DASHBOARD ════════════════════════════════
dash = wb.active
dash.title = "Dashboard"
dash.sheet_view.showGridLines = False
for col, w in (("A", 3), ("B", 26), ("C", 18), ("D", 18), ("E", 18), ("F", 3)):
    dash.column_dimensions[col].width = w

# title banner
dash.merge_cells("B2:E3")
t = dash["B2"]
t.value = "SWITCH  ·  2-DAY ORDER BLITZ"
t.fill = fill(INDIGO); t.font = Font(bold=True, color=WHITE, size=18)
t.alignment = Alignment(horizontal="center", vertical="center")
dash.merge_cells("B4:E4")
sub = dash["B4"]
sub.value = "Goal: %d confirmed orders  ·  shifts of 1–12 hours, billed per hour  ·  no monthly" % GOAL
sub.font = Font(color=GREY, size=10, italic=True)
sub.alignment = Alignment(horizontal="center")

def card(cell, label, formula, big=True, color=INDIGO, fmt="0"):
    """A label above a large formula-driven number."""
    col = cell[0]; row = int(cell[1:])
    lc = dash["%s%d" % (col, row)]
    lc.value = label
    lc.font = Font(bold=True, color=GREY, size=9)
    lc.alignment = Alignment(horizontal="center")
    vc = dash["%s%d" % (col, row + 1)]
    vc.value = formula
    vc.number_format = fmt
    vc.font = Font(bold=True, color=color, size=26 if big else 15)
    vc.alignment = Alignment(horizontal="center", vertical="center")
    dash.row_dimensions[row + 1].height = 34
    for rr in (row, row + 1):
        dash["%s%d" % (col, rr)].fill = fill(INDIGO_L if big else WHITE)

contacted = "=COUNTA(%s)" % cs("G2:G%d" % LAST)
interested = "=COUNTIF(%s,\"Yes\")" % cs("G2:G%d" % LAST)
callbacks = "=COUNTIF(%s,\"Call back\")" % cs("G2:G%d" % LAST)
orders = "=COUNTIF(%s,\"Yes\")" % cs("I2:I%d" % LAST)
pending = "=COUNTIF(%s,\"Pending\")" % cs("I2:I%d" % LAST)
hours = "=SUM(%s)" % cs("H2:H%d" % LAST)
total = "=COUNTA(%s)" % cs("B2:B%d" % LAST)

# headline row of cards
card("C6", "ORDERS CONFIRMED", orders, color=GREEN_TX)
card("D6", "GOAL", "=%d" % GOAL, color=DARK)
card("E6", "REMAINING", "=MAX(0,%d-COUNTIF(%s,\"Yes\"))" % (GOAL, cs("I2:I%d" % LAST)), color=INDIGO)

# progress bar (text) row 9
dash.merge_cells("B9:E9")
pb = dash["B9"]
pb.value = ('=REPT("█",ROUND(MIN(1,COUNTIF(%s,"Yes")/%d)*20,0))'
            '&REPT("░",20-ROUND(MIN(1,COUNTIF(%s,"Yes")/%d)*20,0))'
            '&"  "&TEXT(MIN(1,COUNTIF(%s,"Yes")/%d),"0%%")'
            % (cs("I2:I%d" % LAST), GOAL, cs("I2:I%d" % LAST), GOAL, cs("I2:I%d" % LAST), GOAL))
pb.font = Font(bold=True, color=INDIGO, size=13)
pb.alignment = Alignment(horizontal="center")

# secondary cards row 11
card("B11", "LEADS CONTACTED", "=%s&\" / \"&%s" % (contacted[1:], total[1:]), big=False, color=DARK, fmt="General")
card("C11", "INTERESTED (YES)", interested, big=False, color=GREEN_TX)
card("D11", "CALL BACKS", callbacks, big=False, color="B7791F")
card("E11", "HOURS BOOKED", hours, big=False, color=INDIGO)

dash.merge_cells("B14:E14")
note = dash["B14"]
note.value = "↳ These numbers update automatically as you fill the Call Sheet. Open the 'Call Sheet' tab to start."
note.font = Font(color=GREY, size=10, italic=True)
note.alignment = Alignment(horizontal="center")

# ════════════════════════════ 2. CALL SHEET ═══════════════════════════════
ws = wb.create_sheet(SHEET)
ws.sheet_view.showGridLines = False

for c, (title, width) in enumerate(COLS, start=1):
    cell = ws.cell(row=1, column=c, value=title)
    cell.fill = fill(INDIGO)
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(c)].width = width
ws.row_dimensions[1].height = 34

WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

for r, row in enumerate(rows, start=2):
    values = [
        row.get("Priority", ""),
        row.get("Business Name", ""),
        row.get("Category", ""),
        row.get("Suggested Roles to Pitch", ""),
        row.get("Phone", ""),
        row.get("Address", ""),
        "", "", "", "", "",
    ]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = CENTER if c in (1, 5, 7, 8, 9, 11) else WRAP
        cell.border = BORDER
        if r % 2 == 0 and c not in (7, 8, 9):
            cell.fill = fill(ZEBRA)
    # priority badge colour
    pc = ws.cell(row=r, column=1)
    pc.font = Font(bold=True, color=INDIGO)
    # map hyperlink
    maps = (row.get("Maps") or "").strip()
    if maps:
        m = ws.cell(row=r, column=11, value="Map")
        m.hyperlink = maps
        m.font = Font(color="0563C1", underline="single")

# dropdowns
def dropdown(col_letter, options):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options),
                        allow_blank=True, showDropDown=False)
    dv.promptTitle = "Pick one"
    dv.prompt = "Click the arrow and choose"
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (col_letter, col_letter, LAST))

dropdown("G", ["Yes", "No", "Maybe", "Call back", "No answer", "Wrong number"])
dropdown("H", [str(i) for i in range(1, 13)])
dropdown("I", ["Yes", "No", "Pending"])

# colour the status cells by value
rng = "G2:I%d" % LAST
ws.conditional_formatting.add(rng, FormulaRule(formula=['G2="Yes"'], fill=fill(GREEN), font=Font(color=GREEN_TX, bold=True)))
ws.conditional_formatting.add(rng, FormulaRule(formula=['G2="No"'], fill=fill(RED)))
ws.conditional_formatting.add("G2:G%d" % LAST, FormulaRule(formula=['G2="Call back"'], fill=fill(AMBER)))
# whole-row tint when an order is confirmed
ws.conditional_formatting.add("A2:K%d" % LAST, FormulaRule(formula=['$I2="Yes"'], fill=fill("EAF8EE")))

ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:K%d" % LAST

# ════════════════════════════ 3. HOW TO USE ═══════════════════════════════
how = wb.create_sheet("How to use")
how.sheet_view.showGridLines = False
how.column_dimensions["A"].width = 3
how.column_dimensions["B"].width = 95

guide = [
    ("SWITCH — 2-DAY ORDER BLITZ  ·  HOW TO USE", "title"),
    ("", ""),
    ("THE OFFER (say this, every time)", "head"),
    ("• Shifts of 1–12 hours, billed by the number of hours. NO monthly, NO weekly.", "body"),
    ("• Pay only AFTER the worker reports — no advance.", "body"),
    ("• Instant replacement if the worker doesn't show or isn't a fit.", "body"),
    ("• Verified, local staff delivered in 24 hours.", "body"),
    ("", ""),
    ("HOW TO WORK THE SHEET", "head"),
    ("1. Open the 'Call Sheet' tab. Work top-down — P1 leads churn staff fastest.", "body"),
    ("2. WhatsApp first, then call. After each contact, click the dropdowns:", "body"),
    ("     · Interested?  → Yes / No / Maybe / Call back / No answer / Wrong number", "body"),
    ("     · Hours Needed → 1 to 12 (what they want per shift)", "body"),
    ("     · Order Confirmed? → Yes once they agree to take staff", "body"),
    ("3. Type anything useful in Remarks (role, timing, when to call back).", "body"),
    ("4. The Dashboard tab counts your orders vs the goal automatically.", "body"),
    ("", ""),
    ("DAILY TARGET", "head"),
    ("Day 1: WhatsApp top 60, call every opener, lock requirements → 4–5 orders.", "body"),
    ("Day 2: call back every 'Maybe' with workers ready, ask each Yes for a referral → 5–6 orders.", "body"),
    ("", ""),
    ("Scripts: see outreach-scripts.md in this folder.", "muted"),
]
r = 1
for text, kind in guide:
    cell = how.cell(row=r, column=2, value=text)
    if kind == "title":
        how.merge_cells("B%d:B%d" % (r, r)); cell.font = Font(bold=True, color=WHITE, size=15)
        cell.fill = fill(INDIGO); cell.alignment = Alignment(horizontal="center", vertical="center")
        how.row_dimensions[r].height = 30
    elif kind == "head":
        cell.font = Font(bold=True, color=INDIGO, size=12)
    elif kind == "muted":
        cell.font = Font(italic=True, color=GREY, size=10)
    else:
        cell.font = Font(color=DARK, size=11)
        cell.alignment = Alignment(wrap_text=True)
    r += 1

wb.active = 0   # open on Dashboard
wb.save(OUT)
print("Wrote %s  ·  3 tabs (Dashboard / Call Sheet / How to use)  ·  %d leads" % (OUT, len(rows)))
